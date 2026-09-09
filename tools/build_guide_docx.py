"""Render ``GUIDE.md`` to ``GUIDE.docx`` with the guide's house style.

Conventions (same as the original hand-styled document):

* body Calibri 9.5 pt, headings from the default Word styles (Title / Heading 1 / Heading 2);
* inline ``code`` in Consolas on a light-grey (EEEEEE) run background;
* fenced blocks tagged ``prompt`` → **yellow boxes** (FFF6DD): text you type into Claude Code exactly as written;
* every other fenced block (``powershell``, ``text``, …) → **grey boxes** (F2F2F2): commands or file contents;
* tables with a dark-blue header row (1F3A5F, white bold text);
* a preface ("Prepared …", "How to read …") and a generated Contents list of all level-1/2 headings.

Usage::

    .venv/Scripts/python.exe tools/build_guide_docx.py            # GUIDE.md -> GUIDE.docx
    .venv/Scripts/python.exe tools/build_guide_docx.py in.md out.docx
"""
from __future__ import annotations

import datetime as _dt
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
GREY_BOX, YELLOW_BOX, INLINE, HEADER_FILL = "F2F2F2", "FFF6DD", "EEEEEE", "1F3A5F"
BODY_PT, CODE_PT = 9.5, 9
INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


# ---------------------------------------------------------------------------------------------------------------------
def _shade(element, fill: str) -> None:
    """Attach a solid ``w:shd`` fill to a run/paragraph/cell properties element."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def add_inline(par, text: str, size: float = BODY_PT, bold_all: bool = False, color: RGBColor | None = None) -> None:
    """Write markdown inline formatting (`code`, **bold**, *italic*) into ``par``."""
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 1:
            run = par.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            run.font.size = Pt(size - 0.5)
            _shade(run._element.get_or_add_rPr(), INLINE)
        elif tok.startswith("**") and tok.endswith("**") and len(tok) > 3:
            run = par.add_run(tok[2:-2]); run.bold = True; run.font.size = Pt(size)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 1:
            run = par.add_run(tok[1:-1]); run.italic = True; run.font.size = Pt(size)
        else:
            run = par.add_run(tok); run.font.size = Pt(size)
        if bold_all:
            run.bold = True
        if color is not None:
            run.font.color.rgb = color


def add_box(doc, lines: list[str], fill: str) -> None:
    """One shaded paragraph holding a code/prompt block (line breaks inside, so the shading is continuous)."""
    par = doc.add_paragraph()
    _shade(par._element.get_or_add_pPr(), fill)
    fmt = par.paragraph_format
    fmt.left_indent = Inches(0.12)
    fmt.space_before = Pt(4)
    fmt.space_after = Pt(6)
    for i, line in enumerate(lines):
        run = par.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(CODE_PT)
        if i < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def add_table(doc, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            if r == 0:
                _shade(cell._element.get_or_add_tcPr(), HEADER_FILL)
                add_inline(par, cell_text, size=BODY_PT - 0.5, bold_all=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                add_inline(par, cell_text, size=BODY_PT - 0.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


# ---------------------------------------------------------------------------------------------------------------------
def build(md_path: Path, out_path: Path) -> int:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.75))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.space_after = Pt(4)

    headings = [(ln.count("#", 0, 3), ln.lstrip("#").strip()) for ln in lines if re.match(r"^#{2,3} ", ln)]

    i = 0
    n_boxes = n_tables = 0
    title_done = False
    while i < len(lines):
        ln = lines[i]

        if ln.startswith("# ") and not title_done:  # document title + preface + contents
            doc.add_paragraph(ln[2:].strip(), style="Title")
            p = doc.add_paragraph()
            today = _dt.date.today()
            add_inline(p, f"Prepared {today.day} {today:%B %Y} for Shammunul Islam — companion to the seaice-py kit "
                          "(regenerated from GUIDE.md by tools/build_guide_docx.py).")
            p = doc.add_paragraph()
            add_inline(p, "**How to read this document:** yellow boxes are prompts you type into Claude Code exactly as "
                          "written (a leading `!` means the command runs in your own shell through Claude Code); grey boxes "
                          "are PowerShell commands, file contents or program output.")
            doc.add_paragraph("Contents", style="Heading 1")
            for level, text in headings:
                p = doc.add_paragraph(style="List Bullet" if level == 2 else "List Bullet 2")
                add_inline(p, text, size=BODY_PT - 0.5)
            title_done = True
            i += 1
            continue

        if ln.startswith("```"):  # fenced block
            lang = ln[3:].strip().lower()
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_box(doc, block, YELLOW_BOX if lang == "prompt" else GREY_BOX)
            n_boxes += 1
            i += 1
            continue

        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*-", lines[i + 1]):  # table
            rows = [split_row(ln)]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            width = len(rows[0])
            rows = [r + [""] * (width - len(r)) for r in rows]
            add_table(doc, rows)
            n_tables += 1
            continue

        if ln.startswith("## "):
            doc.add_paragraph(ln[3:].strip(), style="Heading 1")
        elif ln.startswith("### "):
            doc.add_paragraph(ln[4:].strip(), style="Heading 2")
        elif ln.startswith("#### "):
            doc.add_paragraph(ln[5:].strip(), style="Heading 3")
        elif ln.strip() == "---" or not ln.strip():
            pass
        elif ln.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, ln[2:].strip())
            for run in p.runs:
                run.italic = True
        elif re.match(r"^\s*[-*] ", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"^\s*\d+\. ", ln):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\s*\d+\. ", "", ln))
        else:
            add_inline(doc.add_paragraph(), ln.strip())
        i += 1

    doc.core_properties.title = "seaice-py — Step-by-Step Guide"
    doc.core_properties.author = "Shammunul Islam"
    doc.save(out_path)
    print(f"wrote {out_path.name}: {len(doc.paragraphs)} paragraphs, {n_tables} tables, {n_boxes} boxes, "
          f"{len(headings)} headings, {out_path.stat().st_size / 1e3:.0f} kB")
    return 0


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "GUIDE.md"
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "GUIDE.docx"
    sys.exit(build(src, dst))
